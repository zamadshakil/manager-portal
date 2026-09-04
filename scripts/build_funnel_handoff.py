"""Build the operational handoff from reviewed release source, then render with Word on Windows."""
from pathlib import Path
import argparse
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Hierarchia Sales Funnel Operations Handoff.docx"
SOURCES = ["OPERATIONS.md", "linkedin-and-response-drafts.md", "pilot-proposal-template.md", "TECHNICAL-RUNBOOK.md"]

def inline(paragraph, text):
    pattern = r"(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*|`[^`]+`)"
    for piece in re.split(pattern, text):
        link = re.fullmatch(r"\[([^\]]+)\]\((https?://[^)]+)\)", piece)
        if link:
            element = OxmlElement("w:hyperlink")
            element.set(qn("r:id"), paragraph.part.relate_to(link[2], RT.HYPERLINK, is_external=True))
            run = OxmlElement("w:r")
            props = OxmlElement("w:rPr")
            color = OxmlElement("w:color"); color.set(qn("w:val"), "174F3B"); props.append(color)
            underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single"); props.append(underline)
            run.append(props)
            value = OxmlElement("w:t"); value.text = link[1]; run.append(value)
            element.append(run); paragraph._p.append(element)
        else:
            run = paragraph.add_run(piece.strip("*`") if piece.startswith(("**", "`")) else piece)
            if piece.startswith("**"): run.bold = True
            if piece.startswith("`"): run.font.name = "Consolas"; run.font.size = Pt(10)

def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(.72), Inches(.72)
    sec.left_margin, sec.right_margin = Inches(.85), Inches(.85)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.07
    normal.paragraph_format.widow_control = True
    # Word's bundled Title style can carry a decorative bottom border.
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    for name, size in [("Title", 25), ("Subtitle", 11), ("Heading 1", 16), ("Heading 2", 12), ("Heading 3", 11)]:
        style = doc.styles[name]; style.font.name = "Calibri"; style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(13 if name.startswith("Heading") else 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
    code_style = doc.styles.add_style("Runbook code", 1)
    code_style.font.name = "Consolas"; code_style.font.size = Pt(9.5)
    code_style.paragraph_format.space_after = Pt(2)
    code_style.paragraph_format.line_spacing = 1
    footer = sec.footer.paragraphs[0]
    footer.alignment = 2
    run = footer.add_run("Hierarchia operations  |  "); run.font.size = Pt(9)
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    for index, source in enumerate(SOURCES):
        in_code = False
        for line in (ROOT / "docs" / "sales-assets" / source).read_text(encoding="utf-8").splitlines():
            if line.startswith("```"): in_code = not in_code; continue
            if not line.strip(): continue
            if in_code:
                doc.add_paragraph(line, "Runbook code"); continue
            if line.startswith("#"):
                level = len(line) - len(line.lstrip("#"))
                title = re.sub(r"[^\w\s]", " ", line[level:].strip())
                title = re.sub(r"\s+", " ", title)
                if index == 0 and level == 1: doc.add_paragraph(title, "Title")
                else:
                    heading = doc.add_heading(title, level=min(level if index else level - 1, 3))
                    if index and level == 1: heading.paragraph_format.page_break_before = True
                continue
            style = "List Bullet" if line.startswith("- ") else "Normal"
            paragraph = doc.add_paragraph(style=style)
            inline(paragraph, line[2:] if style == "List Bullet" else line)
    doc.core_properties.title = "Hierarchia Sales Funnel Operations Handoff"
    doc.core_properties.subject = "Live release and sales operator instructions"
    doc.core_properties.author = "ZamDev AI"
    doc.core_properties.last_modified_by = "ZamDev AI"
    doc.save(OUTPUT)
    print(OUTPUT)

def render_word():
    import win32com.client
    from pdf2image import convert_from_path
    folder = ROOT / "docs" / "qa" / "funnel-handoff"
    folder.mkdir(parents=True, exist_ok=True)
    pdf = folder / "handoff.pdf"
    app = win32com.client.DispatchEx("Word.Application")
    app.Visible = False; app.DisplayAlerts = 0
    document = None
    try:
        document = app.Documents.Open(str(OUTPUT), ReadOnly=True)
        document.ExportAsFixedFormat(str(pdf), 17)
    finally:
        if document is not None: document.Close(False)
        app.Quit()
    poppler = Path(r"C:\Users\Dell\.cache\codex-runtimes\codex-primary-runtime\dependencies\native\poppler\Library\bin")
    pages = convert_from_path(pdf, dpi=120, poppler_path=str(poppler))
    for i, page in enumerate(pages, 1): page.save(folder / f"page-{i}.png")
    print(f"Rendered {len(pages)} pages to {folder}")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(); parser.add_argument("--render-word", action="store_true")
    args = parser.parse_args(); build()
    if args.render_word: render_word()
