from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Hierarchia Sales Funnel Implementation Plan.docx"

INK = "172033"
MUTED = "5B6472"
ACCENT = "2E5B88"
PALE = "EAF1F7"
LIGHT = "F3F5F7"
WHITE = "FFFFFF"
GREEN = "E6F3EC"
AMBER = "FFF2D8"
RED = "F9E5E5"
BORDER = "C7CED8"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell, inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_hyperlink(paragraph, text: str, url: str, color=ACCENT, underline=True):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), color)
    props.append(color_node)
    if underline:
        underline_node = OxmlElement("w:u")
        underline_node.set(qn("w:val"), "single")
        props.append(underline_node)
    run.append(props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, before, after in (
        ("Title", 30, 0, 10),
        ("Heading 1", 19, 14, 7),
        ("Heading 2", 13, 11, 5),
        ("Heading 3", 10.5, 8, 3),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title_props = styles["Title"].element.get_or_add_pPr()
    title_border = title_props.find(qn("w:pBdr"))
    if title_border is not None:
        title_props.remove(title_border)

    styles["List Bullet"].font.name = "Aptos"
    styles["List Bullet"].font.size = Pt(9.5)
    styles["List Bullet"].paragraph_format.space_after = Pt(3)
    styles["List Number"].font.name = "Aptos"
    styles["List Number"].font.size = Pt(9.5)
    styles["List Number"].paragraph_format.space_after = Pt(3)


def configure_document(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.72)
        section.right_margin = Inches(0.72)
        section.header_distance = Inches(0.25)
        section.footer_distance = Inches(0.25)

        header = section.header.paragraphs[0]
        header.text = "ZamDev AI  |  Hierarchia go-to-market handoff  |  4 September 2026"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string(MUTED)
        add_page_number(section.footer.paragraphs[0])


def add_paragraph(doc: Document, text: str = "", *, bold_prefix: str | None = None, italic=False):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        first.bold = True
        rest = p.add_run(text[len(bold_prefix):])
        rest.italic = italic
    else:
        run = p.add_run(text)
        run.italic = italic
    return p


def add_bullets(doc: Document, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        if isinstance(item, tuple):
            lead, rest = item
            run = p.add_run(lead)
            run.bold = True
            p.add_run(rest)
        else:
            p.add_run(item)


def add_numbered(doc: Document, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(f"{index}.  {item}")


def add_table(doc: Document, headers, rows, widths=None, first_col_bold=False, row_fills=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header_cells = table.rows[0].cells
    set_repeat_table_header(table.rows[0])
    for index, (cell, header) in enumerate(zip(header_cells, headers)):
        set_cell_shading(cell, ACCENT)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            set_cell_width(cell, widths[index])
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(str(header))
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
        run.font.size = Pt(8.5)
    keep_row_together(table.rows[0])

    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        fill = row_fills[row_index] if row_fills else (LIGHT if row_index % 2 else WHITE)
        for index, (cell, value) in enumerate(zip(cells, row_values)):
            set_cell_shading(cell, fill)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                set_cell_width(cell, widths[index])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            run.font.size = Pt(8.3)
            if first_col_bold and index == 0:
                run.bold = True
        keep_row_together(table.rows[-1])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_status_table(doc: Document, rows):
    fills = []
    for row in rows:
        status = row[0]
        fills.append(GREEN if status == "Ready" else AMBER if status == "Conditional" else RED)
    return add_table(
        doc,
        ["Status", "Area", "Evidence", "Sales implication"],
        rows,
        widths=[0.85, 1.35, 2.65, 2.05],
        first_col_bold=True,
        row_fills=fills,
    )


def add_section_break(doc: Document, title: str, kicker: str | None = None):
    if kicker:
        p = doc.add_paragraph()
        p.paragraph_format.page_break_before = True
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(kicker.upper())
        run.bold = True
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(ACCENT)
    heading = doc.add_heading(title, level=1)
    if not kicker:
        heading.paragraph_format.page_break_before = True


def add_script_block(doc: Document, label: str, subject: str | None, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color=BORDER, size="5")
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT)
    set_cell_margins(cell, top=120, start=150, bottom=120, end=150)
    cell.text = ""
    title = cell.paragraphs[0]
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run(label)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(ACCENT)
    if subject:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        sr = p.add_run(f"Subject: {subject}")
        sr.bold = True
    for line in lines:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.add_run(line)
    keep_row_together(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    configure_document(doc)

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(48)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("HIERARCHIA")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string(ACCENT)

    title = doc.add_paragraph(style="Title")
    title.add_run("Sales Funnel\nImplementation Plan")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    sr = subtitle.add_run("Audit-backed positioning, ICPs, outreach, CRM, demo, pilot and deployment handoff")
    sr.font.size = Pt(14)
    sr.font.color.rgb = RGBColor.from_string(MUTED)

    meta = add_table(
        doc,
        ["Prepared for", "Prepared by", "Date", "Decision horizon"],
        [["ZamDev AI sales and delivery team", "ZamDev AI / Codex", "4 September 2026", "First 90 days"]],
        widths=[1.8, 1.6, 1.35, 1.65],
    )
    doc.add_paragraph()
    add_paragraph(
        doc,
        "Purpose. Give a new sales operator one source of truth for whom to target, what to say, how to qualify, how to run a custom demo and paid pilot, and what must be true before the preview or a customer deployment is shared.",
        bold_prefix="Purpose.",
    )
    add_paragraph(
        doc,
        "Working decision. Position Hierarchia as a custom-deployed document operations and AI review system for service teams—not as a generic project-management tool.",
        bold_prefix="Working decision.",
    )
    add_paragraph(
        doc,
        "Status. The audited build is suitable for an isolated showcase. It is not yet approved for production customer data.",
        bold_prefix="Status.",
    )

    doc.add_page_break()
    doc.add_heading("Executive decision", level=1)
    add_paragraph(
        doc,
        "We should launch a narrow, founder-led funnel aimed first at outsourced service teams with recurring document deliverables. The offer is a paid 30-day workflow pilot in an isolated environment. LinkedIn creates trust and category awareness; targeted email creates conversations; a workflow audit converts interest into a custom demo; the pilot proves operational value before a larger deployment.",
    )
    doc.add_heading("The 30-day objective", level=2)
    add_bullets(
        doc,
        [
            ("One segment: ", "outsource-heavy agencies and business-service teams with 15–150 staff."),
            ("One promise: ", "review every team submission against the operating process before it reaches a client."),
            ("One primary CTA: ", "book a 20-minute workflow audit."),
            ("One commercial offer: ", "a scoped 30-day pilot for one team and up to two recurring workflows."),
            ("One system of record: ", "the existing ZamDev AI Brevo workspace, supplemented by manual LinkedIn activity."),
            ("One success threshold: ", "one paid pilot from the first 400 researched contacts, then recalibrate from evidence."),
        ],
    )
    doc.add_heading("Decisions already made", level=2)
    add_table(
        doc,
        ["Decision", "Choice", "Reason"],
        [
            ["Category", "Document operations and AI review", "Matches the product’s strongest differentiated workflow."],
            ["Primary ICP", "Outsourced service teams", "Visible rework, manager review, deadlines and recurring files."],
            ["CRM and email", "Existing Brevo workspace", "Authenticated domain, verified senders and CRM already available."],
            ["LinkedIn", "Founder-led and manual", "Uses the founder’s credible profile while respecting platform rules."],
            ["Preview", "Isolated showcase only", "Protects real customer, staff and document data."],
            ["Hostname", "Use a name within zamdevai.com", "The requested system.ai.com hostname is outside the confirmed Cloudflare zone."],
        ],
        widths=[1.25, 2.0, 3.2],
        first_col_bold=True,
    )

    add_section_break(doc, "Product audit and sales readiness", "Foundation")
    add_paragraph(
        doc,
        "The codebase audit found a substantial role-based operations portal rather than a simple dashboard. Managers assign document-style work; members submit files; AI can validate, score and summarize against standing team rules plus a task brief. The product also covers deadlines, late reasons, missed submissions, reports, messaging, materials, announcements and team-scoped records.",
    )
    doc.add_heading("Readiness snapshot", level=2)
    add_status_table(
        doc,
        [
            ["Ready", "Build", "Next.js 16.3.4 production build and TypeScript checks pass.", "A controlled showcase can be built from the audited branch."],
            ["Ready", "Dependencies", "Production advisories reduced from 55 to zero at audit time.", "Do not claim permanent security; continue monitoring."],
            ["Ready", "Authentication", "Guest manager and member paths are server-only and provisionable.", "Two role views can be offered in an isolated demo."],
            ["Ready", "Sensitive routes", "Predictable production operations credentials now fail closed.", "A directly exploitable configuration risk was removed."],
            ["Ready", "Telemetry", "Auth and data-minimization controls were added.", "Preview logging has a safer default boundary."],
            ["Conditional", "Automated tests", "No unit, integration or end-to-end test suite is present.", "Do not promise production reliability or enterprise readiness."],
            ["Conditional", "Code quality", "408 legacy compiler and TypeScript warnings remain.", "Plan a hardening sprint before customer production."],
            ["Conditional", "Repository governance", "Branch protection and security automation are not enabled.", "Protect releases before accepting customer data."],
            ["Blocked", "Isolated database", "Supabase free quota is already consumed by unrelated projects.", "Do not reuse an existing production-like database for the showcase."],
        ],
    )
    doc.add_heading("What sales may claim", level=2)
    add_bullets(
        doc,
        [
            "Role-based manager, member and administrator workflows.",
            "Document and image submissions with deadlines and audit history.",
            "Configurable AI review rules and per-task briefs.",
            "Custom deployment and workflow configuration for a scoped pilot.",
            "A live showcase using disposable demo data once the isolated database is provisioned.",
        ],
    )
    doc.add_heading("What sales must not claim yet", level=2)
    add_bullets(
        doc,
        [
            "Enterprise-grade, production-ready, SOC 2 compliant, or guaranteed availability.",
            "Unverified speed, cost-savings or accuracy percentages.",
            "Integrations, SSO, audit certifications or retention guarantees not demonstrated in the scoped environment.",
            "Customer outcomes that have not been measured in a real pilot.",
        ],
    )
    p = doc.add_paragraph()
    p.add_run("Full technical evidence: ").bold = True
    p.add_run("docs/CODEBASE_AUDIT_2026-09-04.md in the manager-portal repository.")

    add_section_break(doc, "Category and positioning", "Strategy")
    doc.add_heading("Category statement", level=2)
    add_paragraph(
        doc,
        "Hierarchia is a custom-deployed document operations and AI review system for teams whose work is delivered as files, images or structured submissions. It gives managers one place to assign work, collect deliverables, enforce deadlines, apply process rules and review exceptions before client delivery.",
    )
    doc.add_heading("One-line pitch", level=2)
    add_script_block(
        doc,
        "Primary pitch",
        None,
        ["Hierarchia helps service teams review every submission against their process before it reaches a client—inside a privately deployed manager portal."],
    )
    doc.add_heading("Thirty-second pitch", level=2)
    add_script_block(
        doc,
        "Spoken version",
        None,
        [
            "Most growing service teams run assignments in chat, store files in Drive and review quality through spreadsheets or manager memory.",
            "Hierarchia brings assignment, submission, deadlines, audit history and configurable AI checks into one portal.",
            "We deploy it around one real workflow first, prove whether it reduces review work and rework, then expand only if the numbers justify it.",
        ],
    )
    doc.add_heading("Message hierarchy", level=2)
    add_table(
        doc,
        ["Layer", "Message", "Evidence to show"],
        [
            ["Business outcome", "Fewer preventable errors reach the client.", "Manager exception queue and review history."],
            ["Operational outcome", "One accountable flow from assignment to approved submission.", "Task, deadline, submission and status views."],
            ["AI role", "Apply repeatable checks; keep managers in control of exceptions.", "Rule configuration and explainable review output."],
            ["Deployment", "Start with one private, isolated workflow.", "Dedicated demo stack and pilot implementation plan."],
            ["Commercial", "Buy a measured pilot before a broader rollout.", "Baseline, target metrics and pilot scorecard."],
        ],
        widths=[1.15, 3.0, 2.3],
        first_col_bold=True,
    )

    add_section_break(doc, "Ideal customer profiles", "Market focus")
    add_paragraph(
        doc,
        "The strongest fit is determined by workflow shape, not industry label. We want teams with recurring file-based deliverables, multiple reviewers, visible deadline pressure and enough volume that manager review has become a bottleneck.",
    )
    doc.add_heading("ICP priority matrix", level=2)
    add_table(
        doc,
        ["Priority", "Segment", "Firmographics", "Pain signal", "Best wedge", "Why now"],
        [
            ["1", "Outsourced business-service teams", "15–150 staff; 3+ team leads; remote or offshore", "Client files reviewed manually across chat, Drive and sheets", "One recurring client-deliverable workflow", "Growth exposes rework, inconsistency and manager bottlenecks"],
            ["2", "E-commerce catalog and content operations", "10–100 operators; high SKU/content volume", "Images, spreadsheets and copy need standardized checks", "One marketplace or catalog submission flow", "Throughput and quality both rise with channel expansion"],
            ["3", "Training and certification providers", "10–75 staff; cohorts or corporate programs", "Assignments, deadlines and rubric checks are fragmented", "One course or certification cohort", "Cohort growth makes manual checking expensive"],
            ["Later", "Regulated document operations", "Claims, logistics, compliance or field operations", "Auditability and missing-data checks", "Non-sensitive proof-of-concept", "Longer security and procurement cycle; product hardening required"],
        ],
        widths=[0.65, 1.45, 1.55, 1.55, 1.4, 1.45],
        first_col_bold=True,
    )
    doc.add_heading("Primary ICP in detail", level=2)
    add_bullets(
        doc,
        [
            ("Company: ", "digital/content agencies, bookkeeping or finance BPOs, recruiting/VA firms and document-heavy shared-service teams."),
            ("Scale: ", "15–150 staff, at least three reviewers or team leads, and 50+ recurring submissions per week."),
            ("Current stack: ", "WhatsApp or Slack for assignment, Google Drive for files, Sheets for status and manager memory for QA."),
            ("Pain: ", "late work, missing fields, inconsistent review, unclear accountability, rework and client-facing quality risk."),
            ("Trigger: ", "headcount growth, a missed deadline, a failed client deliverable, a new offshore team, an audit or a new high-volume account."),
            ("Economic buyer: ", "founder, COO or operations director."),
            ("Champion: ", "department manager who reviews submissions daily."),
            ("Blockers: ", "IT/security, finance and the manager who fears another tool rollout."),
        ],
    )
    doc.add_heading("Disqualifiers", level=2)
    add_bullets(
        doc,
        [
            "Fewer than 10 users and low weekly document volume.",
            "The work unit is a software issue or engineering ticket rather than a file or structured submission.",
            "No measurable manager-review or rework burden.",
            "Enterprise requirements for SSO, formal certifications and procurement before any pilot.",
            "The prospect wants unsupervised AI approval of sensitive or regulated work.",
        ],
    )

    add_section_break(doc, "Buyer committee and qualification", "Revenue process")
    doc.add_heading("Buyer map", level=2)
    add_table(
        doc,
        ["Role", "What they care about", "Question to ask", "Proof to show"],
        [
            ["Founder / COO", "Margins, client retention, scalable delivery", "Where does growth currently add management overhead?", "Pilot economics and expansion plan"],
            ["Operations Director", "Throughput, deadlines, accountability", "Which recurring submission causes the most chasing?", "Assignment and exception workflow"],
            ["Team Manager", "Review workload and team compliance", "What do you check in every submission?", "Configurable rule and review views"],
            ["IT / Security", "Isolation, access, data handling", "What data must never enter the pilot?", "Dedicated environment and data boundary"],
            ["Finance / Procurement", "Cost, risk, contract scope", "What outcome would justify expansion?", "Fixed pilot scope and success criteria"],
        ],
        widths=[1.2, 1.7, 2.0, 1.6],
        first_col_bold=True,
    )
    doc.add_heading("Qualification framework", level=2)
    add_table(
        doc,
        ["Dimension", "Qualified evidence", "Disqualifying answer"],
        [
            ["Volume", "50+ submissions per week or clear peak-load burden", "Only occasional files"],
            ["Review", "A manager checks repeatable criteria", "No standard or reviewer"],
            ["Cost", "Time, rework or deadline risk can be estimated", "No meaningful operational consequence"],
            ["Urgency", "Trigger within 90 days", "General curiosity only"],
            ["Authority", "Buyer and workflow owner will attend demo", "No access to the owner"],
            ["Deployment", "Isolated pilot can exclude sensitive data", "Production integration required on day one"],
            ["Success", "Baseline and target metric can be agreed", "No measurable definition of value"],
        ],
        widths=[1.2, 3.2, 2.1],
        first_col_bold=True,
    )
    doc.add_heading("Discovery questions", level=2)
    add_numbered(
        doc,
        [
            "Walk me through one submission from assignment to client delivery.",
            "How many submissions move through that workflow each week?",
            "What does a manager check every time, and where is that checklist stored?",
            "How long does review take, including chasing missing information?",
            "What percentage comes back for rework, and what causes it?",
            "What happens when a deadline is missed or a low-quality file reaches a client?",
            "Which systems hold the assignment, file, decision and audit trail today?",
            "Which data can be safely used in a 30-day isolated pilot?",
            "Who owns the workflow, budget and technical approval?",
            "What result after 30 days would make you expand to another team?",
        ],
    )

    add_section_break(doc, "Offer design and pricing hypotheses", "Commercial model")
    add_paragraph(
        doc,
        "Pricing below is a validation hypothesis, not a published rate card. The first five qualified opportunities should be used to test willingness to pay, implementation effort and support burden.",
    )
    add_table(
        doc,
        ["Offer", "Scope", "Indicative price", "Conversion purpose"],
        [
            ["Workflow audit", "20-minute discovery plus a one-page workflow map", "No charge", "Qualify pain, volume, authority and deployment fit"],
            ["Paid pilot", "30 days; one team; up to two workflows; up to 25 users; isolated environment", "$1,500–$3,000 setup plus infrastructure", "Prove review-time, deadline or rework improvement"],
            ["Team deployment", "Configured workspace, onboarding, support and agreed integrations", "$4,000–$8,000 setup plus $750–$1,500 monthly", "Operationalize a proven workflow"],
            ["Multi-team rollout", "Multiple departments, governance, priority support and custom integration", "From $10,000 setup plus usage/support", "Expand after verified adoption and value"],
        ],
        widths=[1.2, 2.85, 1.6, 1.85],
        first_col_bold=True,
    )
    doc.add_heading("Pilot success scorecard", level=2)
    add_table(
        doc,
        ["Metric", "Baseline method", "Proposed target", "Owner"],
        [
            ["Manager review minutes per submission", "Time 20 representative reviews", "Agree after baseline; do not pre-promise", "Workflow owner"],
            ["First-pass acceptance rate", "Sample previous 30 submissions", "Directional improvement agreed with buyer", "Team manager"],
            ["On-time submission rate", "Use prior four weeks", "Improvement without hidden manual chasing", "Operations lead"],
            ["Missing-field or format exceptions", "Categorize current rework", "Reduce top two exception categories", "Team manager"],
            ["Weekly active users", "Portal activity during pilot", "At least 80% of pilot users active weekly", "ZamDev delivery"],
            ["Expansion decision", "Executive review on day 28", "Clear go, revise or stop decision", "Economic buyer"],
        ],
        widths=[1.65, 1.75, 2.4, 1.25],
        first_col_bold=True,
    )

    add_section_break(doc, "Showcase landing page", "Acquisition")
    doc.add_heading("Hostname decision", level=2)
    add_paragraph(
        doc,
        "The confirmed Cloudflare zone is zamdevai.com. The requested showcase.system.ai.com is not beneath that zone and currently has no DNS record. Unless ZamDev AI separately owns system.ai.com, use showcase.system.zamdevai.com to preserve the intended naming, or use the shorter showcase.zamdevai.com. Do not configure or advertise system.ai.com without verified ownership.",
    )
    doc.add_heading("Page goal", level=2)
    add_paragraph(
        doc,
        "Convert a qualified visitor into either a workflow-audit booking or an isolated product demo. The page should not collect sensitive documents and should not imply that the preview is a production service.",
    )
    doc.add_heading("Exact landing-page copy", level=2)
    add_script_block(
        doc,
        "Hero",
        None,
        [
            "Eyebrow: Custom-deployed document operations",
            "Headline: Review every team submission against your process before it reaches a client.",
            "Subhead: Assign work, collect files, enforce deadlines and apply configurable AI quality checks in one privately deployed manager portal.",
            "Primary CTA: Book a workflow audit",
            "Secondary CTA: Open the live demo",
            "Trust note: Demo uses isolated sample data. No customer files are stored in the showcase.",
        ],
    )
    add_script_block(
        doc,
        "Problem section",
        None,
        [
            "Heading: Your process should not live in manager memory.",
            "Body: When assignments happen in chat, files live in Drive and quality checks sit in spreadsheets, managers spend time chasing work instead of improving delivery. Hierarchia turns that scattered process into one accountable review flow.",
        ],
    )
    add_table(
        doc,
        ["Section", "Copy direction", "Proof asset"],
        [
            ["How it works", "Assign → submit → validate → review exceptions → report", "Five-step product screenshot strip"],
            ["Use cases", "Agencies, back-office operations, e-commerce content, training providers", "Three workflow cards"],
            ["Manager view", "See deadlines, exceptions, late reasons and submissions needing attention", "Manager dashboard screenshot"],
            ["Member view", "Know what is due, the acceptance criteria and what needs correction", "Member task screenshot"],
            ["AI with control", "Rules support the reviewer; people approve exceptions", "Rule configuration plus review result"],
            ["Deployment", "Start isolated with one team; connect systems only after approval", "Simple deployment diagram"],
            ["Offer", "A 30-day paid pilot with baseline and success scorecard", "Pilot scope summary"],
            ["FAQ", "Security boundary, data, AI role, implementation and pricing", "Plain-language answers"],
        ],
        widths=[1.2, 3.6, 1.9],
        first_col_bold=True,
    )
    doc.add_heading("FAQ copy", level=2)
    add_bullets(
        doc,
        [
            ("Is this another project-management tool? ", "No. The initial use case is the document submission and review process: criteria, files, deadlines, exceptions and approval history."),
            ("Does AI approve work automatically? ", "Not by default. AI applies configured checks and summarizes issues; managers remain responsible for approval and escalation."),
            ("Can we use our own process? ", "Yes. The pilot is configured around one or two repeatable workflows rather than forcing a generic template."),
            ("Where is data hosted? ", "The pilot uses an isolated environment agreed during scoping. Customer-specific production hosting is defined after the pilot."),
            ("How long does setup take? ", "Timing is confirmed after the workflow audit. Do not publish a fixed setup-time promise until the first pilots are measured."),
        ],
    )

    add_section_break(doc, "Funnel architecture and CRM", "Execution system")
    doc.add_heading("Funnel stages", level=2)
    add_table(
        doc,
        ["Stage", "Entry", "Required activity", "Exit criterion", "SLA"],
        [
            ["Target", "Account meets ICP hypothesis", "Research workflow, trigger and buyer", "Named contact with relevant observation", "24 hours"],
            ["Contacted", "First relevant touch sent", "Email plus manual LinkedIn context", "Reply, referral or sequence completion", "14 business days"],
            ["Engaged", "Positive reply or substantive interaction", "Clarify problem and route to audit", "Workflow audit booked", "1 business day"],
            ["Discovery booked", "Meeting accepted", "Research and agenda", "Meeting held or rescheduled", "48 hours"],
            ["Qualified", "Volume, pain, owner, urgency and pilot fit confirmed", "Create workflow map and demo plan", "Custom demo date agreed", "3 business days"],
            ["Custom demo", "Prospect workflow modeled", "Show only relevant manager/member paths", "Pilot scope and buyer next step agreed", "5 business days"],
            ["Proposal", "Commercial and technical fit confirmed", "Send fixed-scope pilot proposal", "Signed, lost or dated follow-up", "5 business days"],
            ["Pilot", "Paid scope and environment approved", "Baseline, configure, onboard, measure", "Executive review and expansion decision", "30 days"],
            ["Closed won", "Payment and deployment decision", "Handoff to delivery and success", "Production plan approved", "2 business days"],
            ["Nurture / lost", "No current decision", "Record reason and dated revisit", "New trigger or closed permanently", "Quarterly"],
        ],
        widths=[1.0, 1.45, 1.75, 1.75, 0.85],
        first_col_bold=True,
    )
    doc.add_heading("Monthly planning model", level=2)
    add_table(
        doc,
        ["Metric", "Planning assumption", "Formula / interpretation"],
        [
            ["Researched contacts", "400", "20 per business day"],
            ["Delivered email", "380", "95% delivery assumption"],
            ["Positive replies", "15", "About 4% of delivered; validate after 100"],
            ["Meetings booked", "10", "Two-thirds of positive replies"],
            ["Meetings held", "8", "80% show rate"],
            ["Qualified opportunities", "5", "Problem, urgency, authority and deployment fit"],
            ["Proposals", "2–3", "Only after custom demo and agreed success metric"],
            ["Paid pilots", "1", "Initial model, not a guarantee"],
        ],
        widths=[1.7, 1.45, 3.5],
        first_col_bold=True,
    )
    doc.add_heading("Brevo configuration", level=2)
    add_paragraph(
        doc,
        "Observed account state: the existing ZamDev AI Brevo workspace is accessible through Zamad’s personal Gmail, contains 197 contacts, has 300 plan emails available in the current period, and shows two verified senders—mail@zamdevai.com and no-reply@zamdevai.com. The zamdevai.com sending domain is authenticated and Brevo reports it compliant with major sender requirements. Existing campaigns and contacts mean the workspace must be changed carefully.",
    )
    add_table(
        doc,
        ["Object", "Exact configuration", "Implementation note"],
        [
            ["Primary sender", "Zamad Shakeel at ZamDev AI <mail@zamdevai.com>", "Use a human sender; retain no-reply for transactional notifications."],
            ["Contact list", "Hierarchia — Prospects", "Create only after confirming how existing contacts are segmented."],
            ["Suppression list", "Hierarchia — Do Not Contact", "Never re-import or re-message opted-out contacts."],
            ["Tags", "HIERARCHIA, ICP_AGENCY, ICP_ECOM, ICP_TRAINING, SOURCE_LINKEDIN, SOURCE_REFERRAL", "Use tags for reporting and routing."],
            ["Lifecycle", "Target, Contacted, Engaged, Discovery, Qualified, Demo, Proposal, Pilot, Won, Lost, Nurture", "If plan limits extra pipelines, store lifecycle as a custom attribute."],
            ["Required fields", "Segment, headcount, workflow, volume/week, current tools, pain, trigger, buyer, lawful basis, next step/date", "Do not send until segment, relevance and source are filled."],
            ["Deal stages", "Discovery, Qualified, Custom demo, Proposal, Pilot, Won, Lost", "Avoid overwriting the default pipeline without reviewing current deals."],
            ["Tracking", "UTM source, medium, campaign, content; privacy-aware open tracking", "Clicks and replies matter more than inflated open rates."],
        ],
        widths=[1.25, 3.25, 2.15],
        first_col_bold=True,
    )
    add_paragraph(
        doc,
        "Change-control note. No sender, list, contact, automation, campaign or deal stage was created during this audit. Those actions can affect an existing marketing workspace and should be implemented only after the owner approves the exact segmentation and confirms whether current contacts belong to unrelated campaigns.",
        bold_prefix="Change-control note.",
    )

    add_section_break(doc, "Email infrastructure and compliance", "Outbound")
    doc.add_heading("Operating rule", level=2)
    add_paragraph(
        doc,
        "Use zamadshakil@gmail.com as the administrator and sign-in identity, not as the cold-outreach From address. Send qualified, low-volume outreach through Brevo from the authenticated zamdevai.com domain. Use mail@zamdevai.com with a human display name so prospects can reply; reserve no-reply@zamdevai.com for transactional notices.",
    )
    doc.add_heading("Technical checklist", level=2)
    add_table(
        doc,
        ["Control", "Current state", "Required action"],
        [
            ["Sending domain", "zamdevai.com authenticated in Brevo", "Re-check authentication before first campaign and monthly thereafter"],
            ["SPF / DKIM / DMARC", "Brevo reports sender compliance and DMARC configured", "Confirm DNS record status in Brevo and Cloudflare after any sender change"],
            ["Reply handling", "mail@zamdevai.com verified", "Verify Cloudflare routing delivers replies to the monitored Gmail inbox"],
            ["From identity", "Current sender name is ZamDev AI", "Use Zamad Shakeel at ZamDev AI for human outreach"],
            ["Postal address", "Not verified during audit", "Add a valid physical postal address to every commercial campaign footer"],
            ["Unsubscribe", "Brevo supports campaign unsubscribe", "Use clear one-click unsubscribe for campaigns and honor direct opt-outs"],
            ["Volume", "Existing plan is limited", "Start 10–20 carefully researched contacts per weekday; do not blast"],
            ["Data source", "Not configured", "Record source URL, date, business relevance and lawful basis"],
        ],
        widths=[1.3, 2.55, 2.85],
        first_col_bold=True,
    )
    doc.add_heading("Compliance guardrails", level=2)
    add_bullets(
        doc,
        [
            "Use accurate sender identity and non-deceptive subjects. Clearly identify ZamDev AI.",
            "Include a valid physical postal address and a clear opt-out mechanism in commercial messages.",
            "Honor opt-outs promptly; U.S. CAN-SPAM requires honoring requests within 10 business days.",
            "For UK business outreach, distinguish corporate subscribers from sole traders and partnerships; retain a legitimate-interest assessment where appropriate.",
            "Do not import personal addresses merely because they are discoverable. Target named business roles only when the message is relevant to their work.",
            "Apply the strictest applicable rule for the recipient’s jurisdiction and obtain legal advice before scaling into unfamiliar markets.",
            "This plan is operational guidance, not legal advice.",
        ],
    )

    add_section_break(doc, "Email sequence", "Outbound playbook")
    add_paragraph(
        doc,
        "The sequence is intentionally short and problem-led. Every bracketed field must be researched. Never invent a trigger, claim a case study that does not exist, or imply that a tailored demo has been prepared when it has not.",
    )
    add_script_block(
        doc,
        "Touch 1 — Day 1",
        "quick question about [workflow]",
        [
            "Hi [First name],",
            "I noticed [Company] runs [specific service / workflow]. When [team] submits [document or image type], how much manager time goes into checking completeness, format and deadlines before it reaches a client?",
            "We built Hierarchia to put assignment, submission, process checks and exception review in one privately deployed portal.",
            "Worth mapping one workflow for 20 minutes?",
            "Zamad",
            "P.S. If this is not relevant, reply ‘no’ and I will not follow up.",
        ],
    )
    add_script_block(
        doc,
        "Touch 2 — Day 3",
        "re: [workflow] review",
        [
            "Hi [First name],",
            "The useful starting point is not an AI transformation project. It is one repeatable checklist managers already apply to [document type].",
            "We can model that flow, show which submissions need attention and measure whether review time or rework changes during a 30-day pilot.",
            "Is [workflow] owned by you or someone else on the operations team?",
            "Zamad",
        ],
    )
    add_script_block(
        doc,
        "Touch 3 — Day 6",
        "three numbers for [workflow]",
        [
            "Hi [First name],",
            "If you send me three rough numbers—submissions per week, reviewers involved and minutes per review—I can return a one-page workflow map with the likely bottleneck.",
            "No documents or sensitive data needed.",
            "Useful?",
            "Zamad",
        ],
    )
    add_script_block(
        doc,
        "Touch 4 — Day 10",
        "example: manager exception queue",
        [
            "Hi [First name],",
            "The manager view in Hierarchia is designed around exceptions: late work, missing requirements, failed checks and submissions waiting for a decision.",
            "That keeps the pilot focused on the part of the process that consumes manager attention—not on replacing every tool at once.",
            "Should I send the live showcase when the isolated demo environment is available?",
            "Zamad",
        ],
    )
    add_script_block(
        doc,
        "Touch 5 — Day 14",
        "close the loop?",
        [
            "Hi [First name],",
            "I will close the loop after this. If reducing review chasing around [workflow] is a 2026 priority, I can map the current process and show a custom demo.",
            "If it is not on the roadmap, a quick ‘not now’ is helpful and I will stop here.",
            "Zamad",
        ],
    )
    doc.add_heading("Segment-specific first-line library", level=2)
    add_table(
        doc,
        ["Segment", "Use only after verifying", "Example first line"],
        [
            ["Agency / BPO", "Service, team structure, client deliverable", "I saw that [Company] delivers [service] with a distributed team; how are managers checking [deliverable] before client handoff?"],
            ["E-commerce", "Marketplace, SKU volume, content workflow", "Your team appears to manage listings across [channel]; how are images and catalog files checked before upload?"],
            ["Training", "Program, cohort, assignment type", "I saw the [program] offering; how are assignment deadlines and rubric checks managed across reviewers?"],
            ["Referral", "Referrer permission and problem context", "[Referrer] suggested I speak with you about the review workflow behind [process]."],
        ],
        widths=[1.25, 2.55, 3.2],
        first_col_bold=True,
    )

    add_section_break(doc, "LinkedIn strategy and scripts", "Trust channel")
    doc.add_heading("Current asset", level=2)
    add_paragraph(
        doc,
        "Zamad’s profile already carries useful founder proof: Fractional CTO positioning, 25+ ventures, ZamDev AI, AI agents, automation and SaaS MVPs. The missing piece is a clear operational category for Hierarchia.",
    )
    doc.add_heading("Recommended profile updates", level=2)
    add_script_block(
        doc,
        "Headline",
        None,
        ["Fractional CTO & Founder at ZamDev AI | Custom-deployed AI operations software for document-heavy teams | Helped 25+ ventures ship"]
    )
    add_script_block(
        doc,
        "About opening",
        None,
        [
            "I help growing service teams turn manual, document-heavy operations into production software.",
            "Our current focus is Hierarchia: a privately deployed manager portal for assigning work, collecting deliverables, enforcing deadlines and reviewing submissions against repeatable process rules.",
            "We start with one workflow, one team and measurable operational outcomes—then expand only when the evidence supports it.",
        ],
    )
    add_bullets(
        doc,
        [
            "Featured item 1: the showcase landing page.",
            "Featured item 2: a 90-second manager/member product walkthrough.",
            "Featured item 3: a one-page workflow audit offer.",
            "Banner message: Review every team submission before it reaches a client.",
        ],
    )
    doc.add_heading("Manual outreach sequence", level=2)
    add_script_block(
        doc,
        "Connection note",
        None,
        ["Hi [First] — I’m researching how [segment] teams review [deliverable] before client handoff. Your work at [Company] looked relevant. Open to connecting?"]
    )
    add_script_block(
        doc,
        "After acceptance",
        None,
        ["Thanks, [First]. Quick research question: which part of [workflow] causes more management work—chasing submissions, checking them, or sending rework back?"]
    )
    add_script_block(
        doc,
        "Value follow-up",
        None,
        ["That matches the workflow we are designing Hierarchia around. If you share rough volume, reviewers and review time, I can map the current bottleneck on one page—no sensitive data needed."]
    )
    add_script_block(
        doc,
        "Meeting ask",
        None,
        ["Would a 20-minute workflow audit next week be useful? I’ll show the current-state map first; a product demo only if the fit is real."]
    )
    add_paragraph(
        doc,
        "Platform rule. Keep research, connection requests and messages manual. LinkedIn prohibits unauthorized bots, scraping and automated message or connection activity. Use the CRM to schedule human tasks; do not use it to drive an unauthorized LinkedIn bot.",
        bold_prefix="Platform rule.",
    )
    doc.add_heading("Three-post weekly content system", level=2)
    add_table(
        doc,
        ["Day", "Content pillar", "Prompt", "CTA"],
        [
            ["Tuesday", "Operational pain", "Show how chat + Drive + Sheets creates invisible review work", "Ask readers where their checklist lives"],
            ["Thursday", "Build in public", "Show one Hierarchia workflow or design decision without customer data", "Invite operations leaders to challenge it"],
            ["Saturday", "Founder proof", "Share a lesson from shipping 25+ ventures or a custom deployment", "Offer the workflow audit"],
        ],
        widths=[0.8, 1.35, 3.35, 1.25],
        first_col_bold=True,
    )
    doc.add_heading("Draft post 1", level=2)
    add_script_block(
        doc,
        "Your process should not live in manager memory",
        None,
        [
            "A growing service team usually does not have a ‘work management’ problem.",
            "It has a review problem.",
            "The assignment is in chat. The file is in Drive. The status is in a sheet. The acceptance criteria live in the manager’s head.",
            "Every new hire adds more chasing and more inconsistent QA.",
            "We are building Hierarchia around a narrower question: can every submission be checked against the team’s process before it reaches a client?",
            "Where does your review checklist live today?",
        ],
    )
    doc.add_heading("Draft post 2", level=2)
    add_script_block(
        doc,
        "AI should surface exceptions, not hide decisions",
        None,
        [
            "The least useful AI feature is a score nobody trusts.",
            "For document-heavy operations, the better pattern is simple:",
            "1. Managers define the standing rules. 2. Each task adds a specific brief. 3. AI checks the submission. 4. People decide the exception.",
            "The goal is not unsupervised approval. It is fewer preventable review cycles and a clearer audit trail.",
            "That is the operating model we are testing with Hierarchia.",
        ],
    )
    doc.add_heading("Draft post 3", level=2)
    add_script_block(
        doc,
        "Start with one workflow",
        None,
        [
            "Most software rollouts fail by trying to replace everything at once.",
            "Our pilot approach is deliberately smaller: one team, up to two recurring workflows, a baseline and a 30-day decision.",
            "If manager review time, first-pass acceptance or on-time delivery does not improve, we should know quickly.",
            "If it does, expansion has evidence behind it.",
            "I’m opening a small number of workflow audits for service teams. Message me with the document process that creates the most chasing.",
        ],
    )

    add_section_break(doc, "Discovery, demo, proposal and pilot", "Sales-to-delivery SOP")
    doc.add_heading("Discovery call agenda — 20 minutes", level=2)
    add_table(
        doc,
        ["Time", "Activity", "Output"],
        [
            ["0–2", "Confirm purpose and permission to ask operational questions", "Shared agenda"],
            ["2–8", "Map assignment, submission, review, rework and approval", "Current-state flow"],
            ["8–13", "Quantify volume, review time, rework and deadline impact", "Baseline hypothesis"],
            ["13–16", "Confirm owner, urgency, data boundary and technical constraints", "Qualification decision"],
            ["16–19", "Explain pilot and identify the custom-demo scenario", "Demo brief"],
            ["19–20", "Agree one dated next step", "Calendar commitment or disqualification"],
        ],
        widths=[0.7, 4.2, 1.8],
        first_col_bold=True,
    )
    doc.add_heading("Custom demo — 30 minutes", level=2)
    add_numbered(
        doc,
        [
            "Restate the prospect’s workflow and success metric; obtain confirmation.",
            "Show the manager assigning one realistic task with standing rules and a task brief.",
            "Switch to the member role and submit sample, non-sensitive content.",
            "Return to the manager view and review deadlines, exceptions and AI-supported findings.",
            "Show reporting or audit history only if it relates to the buyer’s stated pain.",
            "Explain the isolation boundary and what the pilot deliberately excludes.",
            "End with pilot scope, owners, baseline method, price and decision date.",
        ],
    )
    doc.add_heading("Proposal structure", level=2)
    add_bullets(
        doc,
        [
            "Problem and current-state workflow, using the buyer’s language.",
            "Pilot scope: one team, up to two workflows, user limit and isolated environment.",
            "Success metrics and baseline method.",
            "Data, security, AI and integration boundaries.",
            "Implementation schedule and named owners.",
            "Fixed pilot price, infrastructure costs and payment schedule.",
            "Expansion options explicitly outside the pilot.",
            "Go, revise or stop decision on day 28–30.",
        ],
    )
    pilot_heading = doc.add_heading("Thirty-day pilot plan", level=2)
    pilot_heading.paragraph_format.page_break_before = True
    add_table(
        doc,
        ["Phase", "Days", "Work", "Exit gate"],
        [
            ["Baseline", "1–3", "Confirm process, data boundary, sample set and baseline metrics", "Signed workflow and success scorecard"],
            ["Configure", "4–8", "Set roles, team, tasks, rules, sample data and branding", "Buyer accepts the demo scenario"],
            ["Onboard", "9–10", "Train manager and pilot users", "All users complete one practice submission"],
            ["Operate", "11–24", "Run live pilot with approved non-sensitive or customer-authorized data", "Weekly metric review"],
            ["Evaluate", "25–28", "Compare baseline, adoption, exceptions and feedback", "Documented result and gaps"],
            ["Decide", "29–30", "Executive review and expansion recommendation", "Go, revise or stop"],
        ],
        widths=[1.0, 0.8, 3.45, 1.55],
        first_col_bold=True,
    )

    add_section_break(doc, "Deployment and showcase controls", "Delivery")
    doc.add_heading("Required architecture", level=2)
    add_table(
        doc,
        ["Layer", "Showcase choice", "Boundary"],
        [
            ["Source", "Audited preview branch", "Do not deploy an older unpatched commit"],
            ["Frontend / app", "Dedicated preview project", "No connection to unrelated Vercel team projects"],
            ["Database / auth", "New disposable Supabase project or approved equivalent", "Never use ApnaTask, ZamDev AI CRM or another existing project"],
            ["Accounts", "Guest manager and guest member", "Rotate credentials and reset demo data periodically"],
            ["Storage", "Demo-only bucket and sample files", "No real customer or employee documents"],
            ["Email / AI", "Disabled unless the demo explicitly needs them", "No live campaigns, provider secrets or production automations"],
            ["DNS", "showcase.system.zamdevai.com or showcase.zamdevai.com", "Cloudflare-managed owned zone only"],
            ["Monitoring", "Minimal privacy-aware error telemetry", "No request bodies, tokens or customer content"],
        ],
        widths=[1.2, 2.65, 2.85],
        first_col_bold=True,
    )
    doc.add_heading("Current blockers", level=2)
    add_bullets(
        doc,
        [
            "Supabase account has no remaining free project slot; an isolated database cannot be created without archiving an existing project, using another approved organization, or selecting a paid/alternative environment.",
            "A prior Vercel deployment attempt reported a team-membership mismatch. The dedicated personal import must be used or the correct team membership resolved.",
            "The requested system.ai.com hostname is not inside the confirmed Cloudflare zone.",
            "The public landing route and production screenshot set still need to be built and reviewed.",
        ],
    )
    doc.add_heading("Go-live checklist", level=2)
    add_bullets(
        doc,
        [
            "Confirm the hostname inside zamdevai.com.",
            "Provision a disposable database with no copied production records.",
            "Run all migrations on an empty database and fail the deployment on migration errors.",
            "Provision guest manager and guest member accounts; verify role isolation.",
            "Load only synthetic demo content and non-sensitive sample files.",
            "Disable or sandbox transactional email, AI, R2 and external automations.",
            "Run build, TypeScript, lint, dependency audit and a manual role-based smoke test.",
            "Add a preview banner stating that data is synthetic and the environment is not production.",
            "Set data-reset and guest-password rotation ownership.",
            "Create Cloudflare DNS only after the deployment URL is healthy.",
        ],
    )

    add_section_break(doc, "Measurement and operating cadence", "Management")
    doc.add_heading("Weekly dashboard", level=2)
    add_table(
        doc,
        ["Layer", "Metric", "Target / rule", "Action if weak"],
        [
            ["Research", "Accounts passing ICP checklist", "80+ per week", "Tighten data source or segment"],
            ["Deliverability", "Delivery rate", "At least 95%", "Pause source and validate addresses"],
            ["Messaging", "Positive reply rate", "Initial hypothesis 3–5%", "Review relevance and first-line quality"],
            ["Conversion", "Booked / positive replies", "At least 60%", "Strengthen CTA and qualification"],
            ["Attendance", "Held / booked", "At least 75%", "Improve confirmation and agenda"],
            ["Qualification", "Qualified / held", "At least 50%", "Improve targeting or discovery"],
            ["Commercial", "Proposal / qualified", "40–60%", "Fix demo-to-pilot alignment"],
            ["Revenue", "Paid pilot / proposal", "Initial hypothesis 33–50%", "Review risk, price and decision process"],
            ["Content", "Qualified conversations from posts", "Track, no vanity target", "Shift topics toward buyer pain"],
        ],
        widths=[1.1, 1.8, 1.9, 2.45],
        first_col_bold=True,
    )
    doc.add_heading("Cadence", level=2)
    add_bullets(
        doc,
        [
            ("Daily: ", "research 20 accounts, send approved low-volume touches, log replies and set next-step dates."),
            ("Monday: ", "review pipeline aging, no-next-step deals, opt-outs and delivery issues."),
            ("Wednesday: ", "listen to discovery calls and score message-market fit."),
            ("Friday: ", "publish the funnel dashboard, record learnings and approve next week’s test."),
            ("After 100 delivered contacts: ", "revisit ICP, first-line patterns and reply rate."),
            ("After 10 discovery calls: ", "revisit pain language, pricing and pilot scope."),
            ("After each pilot: ", "publish an internal outcome review before making any public claim."),
        ],
    )

    add_section_break(doc, "Thirty sixty ninety day implementation", "Handoff")
    add_table(
        doc,
        ["Window", "Outcome", "Work"],
        [
            ["Days 1–10", "Foundation ready", "Approve ICP and offer; select owned hostname; resolve isolated database; build landing page; create CRM fields, lists and suppression process; verify reply routing"],
            ["Days 11–30", "First evidence", "Publish nine founder posts; research 200–400 accounts; start 10–20 daily emails; run workflow audits; deliver custom demos; record objections"],
            ["Days 31–60", "First paid pilot", "Close and launch one pilot; measure baseline; refine onboarding and proposal; collect approved proof only after results"],
            ["Days 61–90", "Repeatable motion", "Narrow winning subsegment; adjust price; document case evidence; expand content and referrals; decide whether paid data or sales tooling is justified"],
        ],
        widths=[1.0, 1.4, 4.45],
        first_col_bold=True,
    )
    doc.add_heading("RACI", level=2)
    add_table(
        doc,
        ["Workstream", "Responsible", "Accountable", "Consulted", "Informed"],
        [
            ["ICP research and list building", "Sales operator", "Zamad", "Operations SME", "Delivery"],
            ["Email copy and sending", "Sales operator", "Zamad", "Legal/privacy as needed", "Delivery"],
            ["LinkedIn content", "Zamad / content operator", "Zamad", "Sales operator", "Delivery"],
            ["Discovery and custom demo", "Zamad", "Zamad", "Product lead", "Sales operator"],
            ["Preview deployment", "Product / DevOps", "Zamad", "Security reviewer", "Sales operator"],
            ["Pilot implementation", "Delivery lead", "Zamad", "Customer workflow owner", "Sales operator"],
            ["Metrics and reporting", "Sales operator", "Zamad", "Delivery lead", "Team"],
        ],
        widths=[1.8, 1.35, 1.1, 1.55, 1.2],
        first_col_bold=True,
    )
    doc.add_heading("Operator’s first-day checklist", level=2)
    add_numbered(
        doc,
        [
            "Read the codebase audit and this plan; write down any unsupported claim you find elsewhere.",
            "Create the ICP checklist and score 25 example accounts before sending anything.",
            "Confirm the valid physical postal address and monitored reply mailbox.",
            "Get owner approval for the Brevo list, lifecycle field and suppression design.",
            "Prepare 20 personalized Touch 1 drafts and submit them for review.",
            "Schedule the first three LinkedIn posts; publish manually after owner approval.",
            "Set a weekly funnel review with Zamad and delivery.",
        ],
    )
    approval_heading = doc.add_heading("Final pre-launch approval", level=2)
    approval_heading.paragraph_format.page_break_before = True
    add_table(
        doc,
        ["Approval", "Must be true", "Owner sign-off"],
        [
            ["Message", "No unverified speed, ROI, security or customer-result claims", "Zamad"],
            ["Compliance", "Sender identity, postal address, opt-out and source record are complete", "Zamad / counsel as needed"],
            ["CRM", "Existing contacts and campaigns will not be disrupted", "Zamad"],
            ["Preview", "Synthetic data, isolated database, guest reset plan and health checks pass", "Product / DevOps"],
            ["Domain", "Hostname is inside a verified owned Cloudflare zone", "Zamad"],
            ["Offer", "Pilot scope, price hypothesis and success scorecard are approved", "Zamad"],
        ],
        widths=[1.1, 4.2, 1.55],
        first_col_bold=True,
    )

    add_section_break(doc, "Sources and evidence", "References")
    add_paragraph(doc, "Internal evidence", bold_prefix="Internal evidence")
    add_bullets(
        doc,
        [
            "manager-portal/docs/CODEBASE_AUDIT_2026-09-04.md",
            "manager-portal/docs/PROJECT_STATUS.md",
            "manager-portal/docs/CLIENT_DELIVERY/01_EXECUTIVE_OVERVIEW.md",
            "Read-only review of the signed-in Brevo sender, domain, CRM and usage screens on 4 September 2026.",
            "Read-only review of the signed-in LinkedIn founder profile on 4 September 2026.",
            "DNS queries for zamdevai.com and candidate showcase hostnames on 4 September 2026.",
        ],
    )
    doc.add_heading("External sources", level=2)
    sources = [
        ("Google Email sender guidelines", "https://support.google.com/mail/answer/81126?hl=en"),
        ("Google Email sender guidelines FAQ", "https://support.google.com/mail/answer/14229414?hl=en"),
        ("Google Workspace Set up SPF", "https://support.google.com/a/answer/33786?hl=en-eu"),
        ("Google Recommended DMARC rollout", "https://support.google.com/a/answer/10032473?hl=en-in"),
        ("U.S. Federal Trade Commission CAN-SPAM compliance guide", "https://www.ftc.gov/business-guidance/resources/can-spam-act-compliance-guide-business"),
        ("UK Information Commissioner’s Office B2B marketing guidance", "https://ico.org.uk/for-organisations/direct-marketing-and-privacy-and-electronic-communications/business-to-business-marketing/"),
        ("LinkedIn Prohibited software and extensions", "https://www.linkedin.com/help/linkedin/answer/a1341387/prohibited-software-and-extensions?lang=en"),
        ("McKinsey The state of AI in 2025", "https://www.mckinsey.com/capabilities/quantumblack/our-insights/the-state-of-ai"),
        ("McKinsey The state of AI and workflow redesign", "https://www.mckinsey.com/capabilities/quantumblack/our-insights/the-state-of-ai-how-organizations-are-rewiring-to-capture-value"),
        ("LinkedIn and Ipsos 2025 B2B Marketing Benchmark", "https://business.linkedin.com/content/dam/business/marketing-solutions/global/en_US/site/pdf/wp/2025-be-category-famous-and-accelerate-social-trust.pdf"),
    ]
    for label, url in sources:
        p = doc.add_paragraph(style="List Bullet")
        add_hyperlink(p, label, url)
    add_paragraph(
        doc,
        "Interpretation note. External research supports the emphasis on trust, workflow redesign, measurable value and sender compliance. Funnel conversion rates, pricing and success targets in this document are internal planning hypotheses and must be recalibrated from ZamDev AI’s own results.",
        bold_prefix="Interpretation note.",
    )

    core = doc.core_properties
    core.title = "Hierarchia Sales Funnel Implementation Plan"
    core.subject = "Audit-backed sales funnel and execution handoff for Hierarchia"
    core.author = "ZamDev AI"
    core.keywords = "Hierarchia, ZamDev AI, ICP, sales funnel, Brevo, LinkedIn, showcase, pilot"

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
